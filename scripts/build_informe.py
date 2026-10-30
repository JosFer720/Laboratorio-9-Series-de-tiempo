"""Genera el informe final en Word y PDF a partir de los artefactos del laboratorio.

Formato: Times New Roman 12, subtítulos en 14 sin negrita, sin negrita en el resto
y un máximo de cinco páginas. Requiere ``python-docx`` y Microsoft Word para la
conversión a PDF.
"""

from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from forecasting_lab.training import PanelTensors, daily_origins  # noqa: E402

ART = ROOT / "artifacts"
OUT = ROOT / "report"
FIGURAS = OUT / "figuras"
REPOSITORIO = "https://github.com/JosFer720/Laboratorio-9-Series-de-tiempo"
FUENTE = "Times New Roman"

NOMBRES = {
    "Naive": "Ingenuo",
    "SeasonalNaive24": "Ingenuo estacional 24 h",
    "SeasonalNaive168": "Ingenuo estacional 168 h",
    "AutoETS24": "AutoETS",
    "LSTM-AR": "LSTM autorregresiva",
    "Seq2Seq": "Seq2seq",
}
FAMILIAS = {
    "Baseline": "Referencia", "StatsForecast": "Estadístico", "Árboles": "Árboles",
    "Fundacional": "Fundacional", "MLP y lineales": "Lineal y MLP", "Recurrentes": "Recurrente",
    "Encoder-decoder": "Cod.-decod.", "Residuales y convolucionales": "Residual",
    "Atención": "Atención", "Probabilísticos": "Probabilístico",
}
ESTRATEGIAS = {
    "multi-salida": "Multisalida", "autorregresiva": "Autorregresiva", "encoder-decoder": "Cod.-decod.",
    "directa": "Directa", "zero-shot": "Multisalida",
}
DESTACADOS_HORIZONTE = ["N-HiTS", "iTransformer", "LSTM", "LSTM-AR", "DLinear", "DeepAR", "Chronos-Bolt",
                        "SeasonalNaive24"]
# Chronos-Bolt small: parámetros contados con BaseChronosPipeline.from_pretrained.
PARAMETROS_CHRONOS = 47_718_016


def nombre(modelo: str) -> str:
    return NOMBRES.get(modelo, modelo)


def numero(valor: float, decimales: int = 3) -> str:
    """Formato en español: coma decimal y espacio como separador de miles."""

    if valor is None or (isinstance(valor, float) and np.isnan(valor)):
        return "–"
    texto = f"{valor:,.{decimales}f}".replace(",", " ").replace(".", ",")
    return texto


# ---------------------------------------------------------------- datos


def cargar_resultados() -> tuple[pd.DataFrame, pd.DataFrame]:
    """Tabla comparativa y MAE por paso para los 17 modelos."""

    redes = pd.read_csv(ART / "resultados_redes.csv")
    otros = pd.read_csv(ART / "resultados_modelos.csv")
    otros.loc[otros["modelo"].eq("Chronos-Bolt"), ["total_parameters", "trainable_parameters"]] = PARAMETROS_CHRONOS
    otros["mase_navidad"] = otros["modelo"].map(mase_navidad_referencias())
    tabla = pd.concat([redes, otros], ignore_index=True).sort_values("mase").reset_index(drop=True)
    horizonte = pd.concat(
        [pd.read_csv(ART / "mae_horizonte_modelos.csv"),
         pd.read_csv(ART / "detalle_redes" / "metricas_horizonte.csv")[["modelo", "horizon", "mae"]]],
        ignore_index=True,
    )
    return tabla, horizonte


def mase_navidad_referencias() -> dict[str, float]:
    """MASE del 22 al 28 de diciembre para los modelos sin desglose guardado."""

    panel = pd.read_csv(ART / "cache" / "panel_seleccionado.csv.gz", index_col="ds", parse_dates=["ds"])
    data = PanelTensors(panel, "cpu")
    origenes = daily_origins("2014-12-22", "2014-12-28")
    escala = pd.DataFrame({o: data.mase_scale(data.position(o)) for o in origenes}, index=data.clients)
    escala = escala.T.stack().rename("mase_scale").rename_axis(["origin", "unique_id"]).reset_index()
    real = panel.stack().rename("y").rename_axis(["ds", "unique_id"]).reset_index()
    resultado = {}
    for archivo in ("pronosticos_baselines.csv", "pronosticos_statsforecast.csv",
                    "pronosticos_lightgbm.csv", "pronosticos_chronos.csv"):
        pron = pd.read_csv(ART / archivo, parse_dates=["ds"])
        pron = pron.loc[(pron["ds"] >= "2014-12-22") & (pron["ds"] < "2014-12-29")]
        pron = pron.merge(real, on=["unique_id", "ds"]).assign(origin=lambda d: d["ds"].dt.floor("D"))
        pron = pron.merge(escala, on=["origin", "unique_id"])
        pron["error"] = (pron["y"] - pron["y_hat"]).abs() / pron["mase_scale"]
        resultado.update(pron.groupby("modelo")["error"].mean().to_dict())
    return resultado


# ---------------------------------------------------------------- figuras


def estilo() -> None:
    plt.rcParams.update({
        "font.family": "serif", "font.serif": [FUENTE], "font.size": 9,
        "axes.titlesize": 9.5, "axes.labelsize": 9, "legend.fontsize": 7,
        "axes.spines.top": False, "axes.spines.right": False,
    })


def figura_pareto(tabla: pd.DataFrame) -> Path:
    """MASE de prueba contra tiempo de entrenamiento, ambos en escala log."""

    fig, eje = plt.subplots(figsize=(6.6, 2.55))
    sin_entrenar = 0.3
    x = tabla["training_seconds"].fillna(0).clip(lower=sin_entrenar)
    y = tabla["mase"]
    techo = 1.8
    eje.axvspan(0.2, 0.45, color="#eeeeee")
    eje.text(0.3, 1.74, "sin\nentrenar", ha="center", va="top", fontsize=7, color="#555555")
    visibles = y < techo
    eje.scatter(x[visibles], y[visibles], s=18, color="#1f4e79", zorder=3)
    orden = np.argsort(x.to_numpy())
    frontera, mejor = [], np.inf
    for i in orden:
        if y.iloc[i] < mejor:
            frontera.append(i)
            mejor = y.iloc[i]
    eje.step(x.iloc[frontera], y.iloc[frontera], where="post", color="#c0504d", lw=1, label="frontera de Pareto")
    # (dx, dy, alineación) en puntos, para que las etiquetas del grupo de redes no se encimen.
    desplazamientos = {
        "Chronos-Bolt": (5, -3, "left"), "SeasonalNaive24": (5, -8, "left"), "SeasonalNaive168": (5, 2, "left"),
        "NLinear": (-4, 2, "right"), "DLinear": (-4, -4, "right"), "MLP": (5, -7, "left"),
        "TCN": (-4, 3, "right"), "N-HiTS": (0, -9, "center"), "Seq2Seq": (3, -8, "left"),
        "PatchTST": (3, 3, "left"), "LSTM": (0, -9, "center"), "iTransformer": (5, -4, "left"),
        "LSTM-AR": (3, 3, "left"),
    }
    for (_, fila), xi in zip(tabla.iterrows(), x):
        if fila["mase"] >= techo:
            continue
        dx, dy, alineacion = desplazamientos.get(fila["modelo"], (4, 2, "left"))
        eje.annotate(nombre(fila["modelo"]), (xi, fila["mase"]), xytext=(dx, dy), ha=alineacion,
                     textcoords="offset points", fontsize=6.5)
    fuera = tabla.loc[~visibles.to_numpy()]
    for _, fila in fuera.iterrows():
        eje.text(0.5, 1.68, f"{nombre(fila['modelo'])}, MASE {numero(fila['mase'], 2)}, fuera de escala",
                 fontsize=6.5, va="top", color="#555555")
    eje.set_xscale("log")
    eje.set_yscale("log")
    eje.set_ylim(0.72, techo)
    eje.set_yticks([0.8, 0.9, 1, 1.2, 1.5])
    eje.set_yticklabels(["0,8", "0,9", "1", "1,2", "1,5"])
    eje.yaxis.set_minor_formatter(matplotlib.ticker.NullFormatter())
    eje.set_xlim(0.2, 8000)
    eje.set_xlabel("Tiempo de entrenamiento en segundos")
    eje.set_ylabel("MASE en prueba")
    eje.legend(frameon=False, loc="upper right")
    eje.grid(alpha=0.25, which="both")
    return guardar(fig, "figura_pareto.png")


def figura_horizonte(horizonte: pd.DataFrame) -> Path:
    """MAE por paso en prueba para todos los modelos y estrategias de la LSTM en validación."""

    fig, (izq, der) = plt.subplots(1, 2, figsize=(6.6, 2.5), gridspec_kw={"width_ratios": [1.3, 1]})
    # tab10 sin su gris, que se confunde con el resto de modelos.
    colores = [c for i, c in enumerate(plt.cm.tab10.colors) if i != 7]
    for modelo, grupo in horizonte.groupby("modelo"):
        if modelo not in DESTACADOS_HORIZONTE:
            izq.plot(grupo["horizon"], grupo["mae"], lw=0.7, color="#c8c8c8", zorder=1)
    izq.plot([], [], lw=0.7, color="#c8c8c8", label="resto de modelos")
    for indice, modelo in enumerate(DESTACADOS_HORIZONTE):
        grupo = horizonte.loc[horizonte["modelo"].eq(modelo)]
        izq.plot(grupo["horizon"], grupo["mae"], lw=1.1, color=colores[indice], label=nombre(modelo), zorder=2)
    izq.set_yscale("log")
    izq.set_yticks([30, 50, 100, 200, 400])
    izq.set_yticklabels(["30", "50", "100", "200", "400"])
    izq.yaxis.set_minor_formatter(matplotlib.ticker.NullFormatter())
    izq.set_title("Prueba, todos los modelos")
    izq.legend(frameon=False, fontsize=6, ncol=3, loc="upper left", bbox_to_anchor=(0, 1.0),
               handlelength=1.5, columnspacing=0.8)
    izq.set_ylim(22, 1500)
    estrategias = pd.read_csv(ART / "detalle_redes" / "mae_horizonte_estrategias_validacion.csv")
    for estrategia, grupo in estrategias.groupby("estrategia"):
        der.plot(grupo["horizon"], grupo["mae"], marker="o", ms=2.5, lw=1,
                 label={"encoder-decoder": "codificador-decodificador"}.get(estrategia, estrategia))
    der.set_title("Validación, estrategias de la LSTM")
    der.legend(frameon=False)
    for eje in (izq, der):
        eje.set_xlabel("Paso del horizonte h")
        eje.grid(alpha=0.25)
    izq.set_ylabel("MAE en kWh")
    return guardar(fig, "figura_horizonte.png")


def figura_historia() -> Path:
    """MASE de validación contra meses de historia, con Chronos como referencia."""

    datos = pd.read_csv(ART / "sensibilidad_historia.csv")
    fig, eje = plt.subplots(figsize=(6.6, 1.9))
    for modelo, grupo in datos.groupby("modelo"):
        if modelo.startswith("Chronos"):
            eje.axhline(grupo["mase"].iloc[0], color="black", ls="--", lw=1, label="Chronos-Bolt, referencia")
            continue
        etiqueta = {"Seq2Seq": "Seq2seq, mejor recurrente"}.get(modelo, nombre(modelo))
        eje.plot(grupo["history_months"], grupo["mase"], marker="o", ms=3.5, lw=1.2, label=etiqueta)
    eje.set_xticks([3, 12, 30])
    eje.set_xlabel("Meses de historia de entrenamiento")
    eje.set_ylabel("MASE en validación")
    eje.grid(alpha=0.25)
    eje.legend(frameon=False, ncol=2, loc="upper right")
    return guardar(fig, "figura_historia.png")


def guardar(fig: plt.Figure, archivo: str) -> Path:
    FIGURAS.mkdir(parents=True, exist_ok=True)
    destino = FIGURAS / archivo
    fig.tight_layout(pad=0.4)
    fig.savefig(destino, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return destino


# ---------------------------------------------------------------- documento


def configurar(documento: Document) -> None:
    seccion = documento.sections[0]
    seccion.page_width, seccion.page_height = Cm(21.59), Cm(27.94)
    seccion.orientation = WD_ORIENT.PORTRAIT
    for lado in ("left_margin", "right_margin"):
        setattr(seccion, lado, Cm(2.0))
    seccion.top_margin = seccion.bottom_margin = Cm(1.8)
    normal = documento.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    formato = normal.paragraph_format
    formato.space_before, formato.space_after = Pt(0), Pt(5)
    formato.line_spacing = 1.0


def sin_cortes(texto: str) -> str:
    """Espacios duros para que Word no separe 25 de %, ni 48 de 817 en un salto de línea."""

    texto = texto.replace(" %", " %")
    return re.sub(r"(?<=\d) (?=\d{3}\b)", " ", texto)


def parrafo(documento: Document, texto: str, tamano: float = 12, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
            antes: float = 0, despues: float = 5):
    p = documento.add_paragraph()
    p.alignment = alineacion
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    corrida = p.add_run(sin_cortes(texto))
    corrida.font.name = FUENTE
    corrida.font.size = Pt(tamano)
    corrida.bold = False
    return p


def subtitulo(documento: Document, texto: str) -> None:
    p = parrafo(documento, texto, tamano=14, alineacion=WD_ALIGN_PARAGRAPH.LEFT, antes=6, despues=3)
    p.paragraph_format.keep_with_next = True


def leyenda(documento: Document, texto: str) -> None:
    es_tabla = texto.startswith("Tabla")
    p = parrafo(documento, texto, alineacion=WD_ALIGN_PARAGRAPH.CENTER, antes=2, despues=2 if es_tabla else 6)
    p.paragraph_format.keep_with_next = es_tabla


def figura(documento: Document, ruta: Path, texto: str, ancho: float = 16.8) -> None:
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=Cm(ancho))
    leyenda(documento, texto)


def tabla_word(documento: Document, encabezados: list[str], filas: list[list[str]], tamano: float,
               anchos: list[float]) -> None:
    tabla = documento.add_table(rows=1, cols=len(encabezados))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.style = "Table Grid"
    for celda, texto in zip(tabla.rows[0].cells, encabezados):
        escribir_celda(celda, texto, tamano, sombreado=True)
    for fila in filas:
        celdas = tabla.add_row().cells
        for celda, texto in zip(celdas, fila):
            escribir_celda(celda, texto, tamano)
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos):
            celda.width = Cm(ancho)
    # Separador corto, una línea vacía completa en 12 puntos ocupa demasiado.
    separador = documento.add_paragraph().paragraph_format
    separador.space_after = Pt(0)
    separador.line_spacing = Pt(4)


def escribir_celda(celda, texto: str, tamano: float, sombreado: bool = False) -> None:
    celda.text = ""
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if texto[:1].isdigit() or texto in {"–", "Sí", "No"} else WD_ALIGN_PARAGRAPH.LEFT
    corrida = p.add_run(texto)
    corrida.font.name = FUENTE
    corrida.font.size = Pt(tamano)
    margenes = OxmlElement("w:tcMar")
    for lado in ("left", "right"):
        nodo = OxmlElement(f"w:{lado}")
        nodo.set(qn("w:w"), "45")
        nodo.set(qn("w:type"), "dxa")
        margenes.append(nodo)
    celda._tc.get_or_add_tcPr().append(margenes)
    if sombreado:
        fondo = OxmlElement("w:shd")
        fondo.set(qn("w:val"), "clear")
        fondo.set(qn("w:fill"), "E7E6E6")
        celda._tc.get_or_add_tcPr().append(fondo)


def filas_comparacion(tabla: pd.DataFrame) -> list[list[str]]:
    filas = []
    for _, r in tabla.iterrows():
        parametros = r["total_parameters"]
        filas.append([
            nombre(r["modelo"]),
            FAMILIAS.get(r["familia"], r["familia"]),
            str(int(r["L"])),
            ESTRATEGIAS.get(r["estrategia"], r["estrategia"]),
            "Sí" if bool(r["usa_covariables"]) else "No",
            numero(parametros, 0) if pd.notna(parametros) else "–",
            numero(r["training_seconds"], 0) if r["training_seconds"] > 0 else "–",
            numero(r["inference_seconds"], 0),
            numero(r["mae"], 1),
            numero(r["rmse"], 0),
            numero(r["wape"], 3),
            numero(r["mase"], 3),
            numero(r["mase_navidad"], 2),
        ])
    return filas


FILAS_ITERACIONES = [
    ["Ventana L", "48, 168 y 336", "0,908, 0,889 y 0,884", "0,945, 0,886 y 0,873"],
    ["Historia en meses", "3, 12 y 30", "1,341, 0,965 y 0,889", "0,886, 0,882 y 0,886"],
    ["Normalización", "global, serie y ventana", "1,459, 0,889 y 0,868", "1,595, 0,886 y 0,880"],
    ["Estado oculto", "32, 64 y 128", "0,948, 0,889 y 0,827", "–"],
    ["Capas y dropout", "1, 2 y 2 con 0,2", "0,889, 0,847 y 0,850", "–"],
    ["Tasa de aprendizaje", "0,0003, 0,001 y 0,003", "0,982, 0,889 y 0,835", "–"],
    ["Tamaño de lote", "256, 1024 y 4096", "0,834, 0,889 y 0,994", "–"],
]


def construir_documento(tabla: pd.DataFrame, figuras: dict[str, Path], texto: dict[str, list[str]]) -> Document:
    documento = Document()
    configurar(documento)
    izquierda, centro, derecha = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT
    for linea, alineacion in [
        ("Universidad del Valle de Guatemala", izquierda), ("CC3092 Deep Learning y Sistemas Inteligentes", izquierda),
        ("", izquierda), ("", izquierda),
        ("Pronóstico del consumo eléctrico horario con redes neuronales", centro),
        ("", izquierda), ("", izquierda),
        ("Fabián Morales", derecha), ("Fernando Ruiz", derecha),
    ]:
        parrafo(documento, linea, alineacion=alineacion, despues=0)
    parrafo(documento, "", despues=0)
    for bloque in texto["intro"]:
        parrafo(documento, bloque)
    for clave, titulo in SECCIONES:
        subtitulo(documento, titulo)
        for bloque in texto[clave]:
            if bloque == "[TABLA_ITERACIONES]":
                leyenda(documento, "Tabla 1. Iteraciones de hiperparámetros por factor")
                tabla_word(documento, ["Factor", "Valores", "MASE LSTM", "MASE DLinear"], FILAS_ITERACIONES, 12,
                           [3.9, 4.6, 4.3, 4.3])
            elif bloque == "[TABLA_COMPARACION]":
                # En 12 puntos las 13 columnas no caben en una tabla, así que se parte en dos
                # con el mismo orden por MASE.
                filas = filas_comparacion(tabla)
                leyenda(documento, "Tabla 2. Características de cada modelo")
                tabla_word(documento, ["Modelo", "Familia", "L", "Estrategia", "Cov.", "Parámetros"],
                           [fila[:6] for fila in filas], 12, [4.9, 2.9, 1.1, 3.0, 1.2, 2.5])
                leyenda(documento, "Tabla 3. Costo y error en prueba")
                tabla_word(documento, ["Modelo", "Entren. s", "Infer. s", "MAE", "RMSE", "WAPE", "MASE",
                                       "MASE Navidad"],
                           [fila[:1] + fila[6:] for fila in filas], 12, [4.9, 1.7, 1.6, 1.5, 1.5, 1.5, 1.5, 3.0])
            elif bloque.startswith("[FIGURA_"):
                clave_figura = bloque.strip("[]")
                figura(documento, figuras[clave_figura], LEYENDAS[clave_figura])
            else:
                parrafo(documento, bloque)
    parrafo(documento, f"Repositorio: {REPOSITORIO}", alineacion=WD_ALIGN_PARAGRAPH.LEFT, antes=6)
    return documento


SECCIONES = [
    ("datos", "1. Datos y preparación de la serie"),
    ("investigacion", "2. Investigación de herramientas"),
    ("protocolo", "3. Protocolo de validación y métricas"),
    ("modelos", "4. Modelos y entrenamiento"),
    ("resultados", "5. Comparación de resultados"),
    ("discusion", "6. Discusión"),
    ("conclusiones", "7. Conclusiones"),
]
LEYENDAS = {
    "FIGURA_PARETO": "Figura 1. Frontera de Pareto del entrenamiento",
    "FIGURA_HORIZONTE": "Figura 2. Error por paso del horizonte",
    "FIGURA_HISTORIA": "Figura 3. Sensibilidad a la historia",
}


def leer_texto(ruta: Path) -> dict[str, list[str]]:
    """Lee el texto del informe: bloques ``## clave`` separados por líneas en blanco."""

    secciones: dict[str, list[str]] = {}
    actual = None
    for bloque in re.split(r"\n\s*\n", ruta.read_text(encoding="utf-8")):
        bloque = bloque.strip()
        if not bloque or bloque.startswith("<!--"):
            continue
        if bloque.startswith("## "):
            actual = bloque[3:].strip()
            secciones[actual] = []
            continue
        secciones[actual].append(" ".join(linea.strip() for linea in bloque.splitlines()))
    return secciones


def validar_texto(texto: dict[str, list[str]]) -> None:
    """Reglas de estilo pedidas para el informe."""

    completo = " ".join(" ".join(bloques) for bloques in texto.values())
    if ";" in completo:
        raise ValueError("El informe no debe usar punto y coma")
    if "**" in completo:
        raise ValueError("El informe no debe usar negrita")
    if len(texto["conclusiones"]) != 3:
        raise ValueError("Se esperan exactamente tres conclusiones")
    for leyenda_texto in LEYENDAS.values():
        if len(leyenda_texto.split(". ", 1)[1].split()) > 5:
            raise ValueError(f"Título de figura demasiado largo: {leyenda_texto}")


def a_pdf(docx: Path, pdf: Path) -> int:
    """Convierte con Word y devuelve el número de páginas."""

    comando = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{docx}'); $n = $d.ComputeStatistics(2); "
        f"$d.SaveAs([ref]'{pdf}', [ref]17); $d.Close([ref]0); $w.Quit(); Write-Output $n"
    )
    salida = subprocess.run(["powershell", "-NoProfile", "-Command", comando], capture_output=True, text=True, check=True)
    return int(salida.stdout.strip().splitlines()[-1])


def main() -> None:
    estilo()
    tabla, horizonte = cargar_resultados()
    figuras = {
        "FIGURA_PARETO": figura_pareto(tabla),
        "FIGURA_HORIZONTE": figura_horizonte(horizonte),
        "FIGURA_HISTORIA": figura_historia(),
    }
    texto = leer_texto(OUT / "informe_texto.md")
    validar_texto(texto)
    documento = construir_documento(tabla, figuras, texto)
    docx = OUT / "informe.docx"
    documento.save(docx)
    paginas = a_pdf(docx.resolve(), (OUT / "informe.pdf").resolve())
    print(f"{docx.name} y informe.pdf generados con {paginas} páginas")
    if paginas > 5:
        raise SystemExit("El informe excede cinco páginas")


if __name__ == "__main__":
    main()
